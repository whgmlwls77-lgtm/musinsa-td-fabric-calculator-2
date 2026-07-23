# -*- coding: utf-8 -*-
"""
test_yochuk_report.py
---------------------
요척 산출 PDF(v33) 검증 — Task #38 [2] + Task #38-b 정정 (사장님 확정 2026-07-23).

Task #38-b 핵심:
  1. 마카 배치도 = sparrow 원본 SVG 그대로 cairosvg 렌더 (matplotlib 재렌더링 폐기)
  2. 스타일 요약: 샘플/계산 사이즈 → '사이즈' 1개 통합, 엔진 필드 삭제
  3. 재질별 요척 표: 합계 행 삭제
  4. 재질별 상세: 배치 실패 · 처리 시간 필드 삭제

검증:
  A. _marker_png_for_material: SVG 있으면 PNG, SVG 없으면 None (matplotlib/pieces 아님)
  B. _v33_summary_rows: 사이즈 1개 · 엔진/샘플사이즈/계산 사이즈 필드 없음
  C. _v33_overview_rows: 합계 행 없음
  D. _v33_material_detail_rows: 배치 실패 · 처리 시간 필드 없음
  E. generate_pdf_report_v33: SVG 배치도 이미지 임베드 (XObject)
"""
from __future__ import annotations

import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import app
import svg_render

# cairosvg 가 렌더 가능한 최소 SVG (path + rect).
MIN_SVG = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 50" '
           'width="100%" height="100%">'
           '<rect x="0" y="0" width="100" height="50" fill="#dddddd"/>'
           '<path d="M10,10 L40,10 L40,40 L10,40 z" fill="#888888"/></svg>')


def _row():
    return {"material": "주원단", "fabric_width_cm": 150.0, "marker_length_cm": 60.0,
            "marker_length_yd": 0.66, "efficiency_pct": 80.0, "pieces_count": 3,
            "mirror_count": 0, "total_garments": 1, "runtime_actual_sec": 1.5}


class TestMarkerPngFromSvg(unittest.TestCase):
    """A. 배치도 = sparrow 원본 SVG 렌더 (matplotlib 폐기 확인)."""

    def test_svg_humanized_preferred(self):
        png = app._marker_png_for_material({"svg_content_humanized": MIN_SVG,
                                            "svg_content": ""})
        self.assertIsNotNone(png)
        self.assertTrue(png.startswith(b"\x89PNG"))

    def test_svg_content_fallback(self):
        png = app._marker_png_for_material({"svg_content": MIN_SVG})
        self.assertIsNotNone(png)
        self.assertTrue(png.startswith(b"\x89PNG"))

    def test_no_svg_returns_none(self):
        """SVG 없으면 None — pieces/matplotlib 로 재렌더링하지 않음 (Task #38-b)."""
        self.assertIsNone(app._marker_png_for_material(
            {"pieces_used": [{"piece_id": "p1"}], "placements": []}))

    def test_error_res_returns_none(self):
        self.assertIsNone(app._marker_png_for_material({"error": "배치 실패",
                                                        "svg_content": MIN_SVG}))


class TestV33SummaryRows(unittest.TestCase):
    """B. 스타일 요약 — 사이즈 1개 · 엔진 삭제."""

    def test_labels_merged_and_no_engine(self):
        rows = app._v33_summary_rows({"style": "S", "selected_sizes": ["L"],
                                      "sample_size": "L", "file_name": "f.dxf",
                                      "generated_at": "2026-07-23"})
        labels = [r[0] for r in rows]
        self.assertEqual(labels, ["스타일", "사이즈", "파일", "작성일"])
        self.assertNotIn("엔진", labels)
        self.assertNotIn("샘플사이즈", labels)
        self.assertNotIn("계산 사이즈", labels)

    def test_size_uses_selected(self):
        rows = app._v33_summary_rows({"selected_sizes": ["M", "L"]})
        size_val = dict(rows)["사이즈"]
        self.assertEqual(size_val, "M, L")


class TestV33OverviewRows(unittest.TestCase):
    """C. 재질별 요척 표 — 합계 행 없음."""

    def test_no_total_row(self):
        rows = app._v33_overview_rows([_row(), dict(_row(), material="안감")])
        materials = [r[0] for r in rows[1:]]  # 헤더 제외
        self.assertNotIn("합계", materials)
        # 재질 수 == 데이터 행 수 (합계 없음)
        self.assertEqual(len(rows) - 1, 2)


class TestV33DetailRows(unittest.TestCase):
    """D. 재질별 상세 — 배치 실패 · 처리 시간 없음."""

    def test_no_fail_no_runtime(self):
        labels = [r[0] for r in app._v33_material_detail_rows(_row())]
        self.assertNotIn("배치 실패", labels)
        self.assertNotIn("처리 시간", labels)
        for expect in ("원단 폭", "마카 길이", "효율", "패턴 갯수 (1벌당)",
                       "마카 갯수 (전체)", "마카 벌수"):
            self.assertIn(expect, labels)


class TestKoreanFontFix(unittest.TestCase):
    """Task #38-c: 배치도 한글 tofu(□) 수정 — 렌더 시 실재 한글 폰트 주입.

    PDF 배치도는 래스터 이미지(SVG→PNG)라 한글이 이미지 픽셀로 박힘 →
    PDF 텍스트 추출 불가. 따라서 '폰트 주입 기전'을 결정적으로 검증한다
    (최종 tofu-free 육안 확인은 사장님 시각 검증).
    """

    KR_SVG = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 300 40" '
              'width="300" height="40">'
              '<text x="10" y="25" font-size="16" font-family="sans-serif">'
              '원단 폭 147.3 cm 마카 길이</text></svg>')

    def test_resolve_returns_real_korean_font(self):
        """시스템에 실재하는 한글 폰트를 fc-match 로 확정 (후보 목록 내)."""
        fam = svg_render.resolve_korean_family()
        self.assertIsNotNone(fam, "실재 한글 폰트 미탐지 — fc-match/폰트 환경 확인")
        self.assertIn(fam, svg_render._KR_FONT_CANDIDATES)
        self.assertTrue(svg_render._font_exists(fam))

    def test_inject_replaces_sans_serif(self):
        """렌더 SVG 의 font-family 가 실재 한글 폰트로 치환 (sans-serif 제거)."""
        fam = svg_render.resolve_korean_family()
        injected = svg_render._inject_korean_font(self.KR_SVG, fam)
        self.assertIn(f'font-family="{fam}"', injected)
        self.assertNotIn('font-family="sans-serif"', injected)

    def test_original_svg_not_mutated(self):
        """사장님 절대 준수: sparrow 원본 SVG 자체 수정 X (복사본만 치환)."""
        fam = svg_render.resolve_korean_family()
        _ = svg_render._inject_korean_font(self.KR_SVG, fam)
        self.assertIn('font-family="sans-serif"', self.KR_SVG)  # 원본 그대로

    def test_render_korean_svg_ok(self):
        """한글 SVG → PNG 정상 산출 (렌더 파이프라인 스모크)."""
        png = svg_render.svg_to_png(self.KR_SVG, output_width=600)
        self.assertIsNotNone(png)
        self.assertTrue(png.startswith(b"\x89PNG"))

    def test_font_family_candidates_are_korean(self):
        """후보 목록은 한글 대응 폰트만 (라틴 폴백 방지 — 콤마 나열 tofu 회피)."""
        expect = {"Apple SD Gothic Neo", "NanumGothic", "Noto Sans CJK KR", "Noto Sans KR"}
        self.assertTrue(expect.issubset(set(svg_render._KR_FONT_CANDIDATES)))


class TestV33Excel(unittest.TestCase):
    """Task #38-d: 요척 Excel = 원단별 개별 시트 + 자동 산출 필드 + 배치도 이미지."""

    @staticmethod
    def _ctx(mats):
        import io
        by, srows, order = {}, [], []
        for mat in mats:
            by[mat] = {"svg_content_humanized": MIN_SVG,
                       "placements": [{"piece_id": f"p{i}"} for i in range(6)],
                       "unplaced": []}
            srows.append({"material": mat, "fabric_width_cm": 147.3, "marker_length_cm": 110.0,
                          "marker_length_yd": 1.2, "efficiency_pct": 78.5, "pieces_count": 6,
                          "mirror_count": 0, "total_garments": 1, "runtime_actual_sec": 2.0,
                          "cm_per_garment": 110.0, "yards_per_garment": 1.2})
            order.append(mat)
        return {"nest_all": {"summary_rows": srows, "by_material": by, "materials_ordered": order},
                "style": "MMAPS003_자켓", "file_name": "MMAPS003.dxf",
                "selected_sizes": ["L"], "generated_at": "2026-07-23 16:00"}

    def _wb(self, mats):
        import io
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(app.generate_excel_report_v33(self._ctx(mats))))

    def test_one_sheet_per_material_named_by_code(self):
        wb = self._wb(["주원단", "안감"])
        self.assertEqual(wb.sheetnames, ["SELF", "LINING"])

    def test_used_materials_only(self):
        """사용된 재질만 시트 (POCKETING 없으면 시트도 없음)."""
        wb = self._wb(["주원단"])
        self.assertEqual(wb.sheetnames, ["SELF"])

    def test_yochuk_fields_present(self):
        ws = self._wb(["주원단"])["SELF"]
        labels = {ws.cell(row=r, column=1).value for r in range(1, 25)}
        for expect in ("원단 폭", "마카 길이 (원단 소요 길이)", "원단별 조각 수 (사용 / 전체)",
                       "원단 효율", "전체 면적", "로스 %", "요척 (로스 반영)",
                       "사이즈 갯수 / 비율"):
            self.assertIn(expect, labels)

    def test_marker_image_present(self):
        ws = self._wb(["주원단"])["SELF"]
        self.assertGreaterEqual(len(ws._images), 1)  # 배치도 이미지 삽입

    def test_landscape_page_setup(self):
        ws = self._wb(["주원단"])["SELF"]
        self.assertEqual(ws.page_setup.orientation, "landscape")

    def test_deleted_fields_absent(self):
        """사장님 확정 삭제: 축율·협력사·담당자·재질 코드 컬럼·코멘트·제출일 없음."""
        ws = self._wb(["주원단"])["SELF"]
        allvals = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
        for banned in ("축율", "협력사", "담당자", "코멘트", "제출", "재질 코드", "세로 축율", "가로 축율"):
            self.assertNotIn(banned, allvals, f"삭제 대상 '{banned}' 잔존")


class TestMarkerLayoutUnified(unittest.TestCase):
    """Task #38-e: 원단별 배치도 = 고정 영역 max-fit → 페이지 레이아웃 통일."""

    WIDE_SVG = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 60" '
                'width="100%" height="100%"><rect width="200" height="60" fill="#ccc"/></svg>')
    TALL_SVG = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 60 200" '
                'width="100%" height="100%"><rect width="60" height="200" fill="#ccc"/></svg>')

    def test_fit_into_wide_fills_width(self):
        fw, fh = app._fit_into(2000, 700, 515.0, 340.0)
        self.assertAlmostEqual(fw, 515.0, places=1)   # 가로형 → 폭 딱 맞음
        self.assertLessEqual(fh, 340.0 + 0.1)

    def test_fit_into_tall_fills_height(self):
        fw, fh = app._fit_into(700, 2000, 515.0, 340.0)
        self.assertAlmostEqual(fh, 340.0, places=1)   # 세로형 → 높이 딱 맞음
        self.assertLessEqual(fw, 515.0 + 0.1)

    def test_pdf_container_same_size_any_aspect(self):
        """가로형/세로형 배치도 → 동일 크기 고정 영역 (페이지 레이아웃 통일)."""
        import svg_render
        pw = svg_render.svg_to_png(self.WIDE_SVG, output_width=800)
        ph = svg_render.svg_to_png(self.TALL_SVG, output_width=300)
        tw = app._rl_marker_fixed_area(pw)
        th = app._rl_marker_fixed_area(ph)
        self.assertEqual(tw._colWidths, th._colWidths)
        self.assertEqual(tw._rowHeights, th._rowHeights)
        self.assertEqual(tw._colWidths, [app._PDF_MARKER_AREA_W])
        self.assertEqual(tw._rowHeights, [app._PDF_MARKER_AREA_H])

    def test_excel_image_bounded_same_box(self):
        """가로형/세로형 시트 이미지 모두 동일 박스 안 (한 변 max)."""
        wide = app._excel_marker_image({"svg_content_humanized": self.WIDE_SVG})
        tall = app._excel_marker_image({"svg_content_humanized": self.TALL_SVG})
        for img in (wide, tall):
            self.assertIsNotNone(img)
            self.assertLessEqual(img.width, app._XLSX_MARKER_W + 0.5)
            self.assertLessEqual(img.height, app._XLSX_MARKER_H + 0.5)
        # 가로형은 폭이, 세로형은 높이가 박스에 참 (일관 스케일)
        self.assertAlmostEqual(wide.width, app._XLSX_MARKER_W, delta=0.5)
        self.assertAlmostEqual(tall.height, app._XLSX_MARKER_H, delta=0.5)


class TestPdfMarkerEmbed(unittest.TestCase):
    """E. v33 PDF 에 SVG 배치도 이미지 임베드."""

    def test_v33_pdf_embeds_svg_image(self):
        nest_all = {
            "summary_rows": [_row()],
            "by_material": {"주원단": {"svg_content_humanized": MIN_SVG}},
            "materials_ordered": ["주원단"],
            "unclassified_warnings": [],
            "total_runtime_sec": 1.5,
        }
        context = {
            "nest_all": nest_all, "style": "TESTSTYLE", "file_name": "f.dxf",
            "sample_size": "L", "selected_sizes": ["L"],
            "generated_at": "2026-07-23 10:00", "runtime_seconds_used": 30,
        }
        pdf = app.generate_pdf_report_v33(context)
        self.assertTrue(pdf.startswith(b"%PDF-"))
        self.assertIn(b"/XObject", pdf)  # 배치도 이미지 임베드 증거


if __name__ == "__main__":
    unittest.main(verbosity=2)
