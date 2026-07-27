# -*- coding: utf-8 -*-
"""
test_axis_report.py
-------------------
축율 검증 다운로드 (PDF/Excel) 생성 검증 (Task #38-g 재설계 — 사장님 확정 2026-07-26).

신규 계약 (면적 우선 판정 + 4변 근거 + 정합성 + 오버랩 이미지):
  pairs [{block_name, material, verdict_label, verdict_severity,
          pp_area_mm2, main_area_mm2, area_exp, v_actual, h_actual,
          consistency_flag, consistency_msg, method,
          edges {left/right/top/bottom: (pp,main,exp)}, diameters, corners_main,
          symmetry, overlap_png}]

검증:
  1. build_axis_pdf → 유효 PDF (%PDF-) + 오버랩 이미지 XObject 임베드
  2. build_axis_pdf 빈 조각 안전
  3. build_axis_xlsx → 시트 4개 (요약/원단별요약/조각별상세/오버랩이미지)
  4. 조각별상세 = 면적 우선 컬럼 (판정·면적·세로/가로·정합성·4변)
  5. 오버랩 이미지 시트 임베드
"""
from __future__ import annotations

import io
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from axis_report_pdf import build_axis_pdf
from axis_report_xlsx import build_axis_xlsx


def _overlap_png():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(2.5, 2.5))
    ax.plot([0, 1, 1, 0, 0], [0, 0, 1, 1, 0])
    buf = io.BytesIO()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


def _pair(block="BLK_1_1", sev="critical", area_exp=-6.13, with_png=True):
    return {
        "block_name": block, "material": "SELF", "confidence": "name_shape",
        "verdict_label": "🚨" if sev == "critical" else ("⚠️" if sev == "warning" else "✅"),
        "verdict_severity": sev,
        "pp_area_mm2": 100893.0, "main_area_mm2": 94709.0,
        "area_exp": area_exp, "v_actual": -5.31, "h_actual": -1.09,
        "consistency_flag": False, "consistency_msg": "정합 — 이론 ≈ 실제",
        "method": "corner",
        "edges": {"left": (40.88, 38.52, -5.76), "right": (40.80, 38.82, -4.86),
                  "top": (20.20, 20.17, -0.12), "bottom": (27.49, 26.92, -2.07)},
        "diameters": None, "corners_main": {"LT": (0, 0), "RT": (1, 0), "LB": (0, 1), "RB": (1, 1)},
        "symmetry": [], "overlap_png": _overlap_png() if with_png else None,
    }


def _ctx(pairs=None, material_summary=None):
    return {
        "style": "MWEKS9D02", "file_name_pp": "QC.dxf", "file_name_main": "PP.dxf",
        "target_size": "S", "material": "SELF", "generated_at": "2026-07-27 10:00",
        "declared": {"vertical": 0.0, "horizontal": 0.0},
        "scoped_pp_total": 1800.0, "scoped_main_total": 1690.0, "scoped_exp": -6.11,
        "n_warn": 0, "n_crit": 1,
        "material_summary": material_summary if material_summary is not None else [
            {"material": "SELF", "n_pieces": 1, "n_warn": 0, "n_crit": 1, "worst_sev": "critical"}
        ],
        "pairs": pairs if pairs is not None else [_pair()],
    }


class TestAxisPdf(unittest.TestCase):

    def test_pdf_valid_bytes(self):
        pdf = build_axis_pdf(_ctx())
        self.assertTrue(pdf.startswith(b"%PDF-"))
        self.assertGreater(len(pdf), 1000)

    def test_pdf_overlap_image_embedded(self):
        """오버랩 이미지 있으면 XObject 임베드."""
        pdf = build_axis_pdf(_ctx())
        self.assertIn(b"/XObject", pdf)

    def test_pdf_empty_pairs_safe(self):
        pdf = build_axis_pdf(_ctx(pairs=[], material_summary=[]))
        self.assertTrue(pdf.startswith(b"%PDF-"))

    def test_pdf_no_overlap_safe(self):
        pdf = build_axis_pdf(_ctx(pairs=[_pair(with_png=False)]))
        self.assertTrue(pdf.startswith(b"%PDF-"))


class TestAxisXlsx(unittest.TestCase):

    def _wb(self, **kw):
        return load_workbook(io.BytesIO(build_axis_xlsx(_ctx(**kw))))

    def test_four_sheets(self):
        wb = self._wb()
        self.assertEqual(wb.sheetnames, ["요약", "원단별요약", "조각별상세", "오버랩이미지"])

    def test_detail_header_area_first(self):
        ws = self._wb()["조각별상세"]
        header = [ws.cell(row=1, column=c).value for c in range(1, 15)]
        self.assertEqual(header, ["조각명", "원단", "판정", "PP 면적(mm²)", "메인 면적(mm²)",
                                  "면적%", "세로%", "가로%", "정합성",
                                  "좌변%", "우변%", "상변%", "하변%", "대칭성"])

    def test_no_gf_columns(self):
        """G-F/곡선상단 컬럼 폐기 확인 (사장님 지시 2026-07-27)."""
        ws = self._wb()["조각별상세"]
        header = [ws.cell(row=1, column=c).value for c in range(1, 20)]
        self.assertNotIn("곡선상단", header)
        self.assertNotIn("G-F%", header)

    def test_detail_row_values(self):
        ws = self._wb()["조각별상세"]
        self.assertEqual(ws.cell(row=2, column=1).value, "BLK_1_1")
        self.assertEqual(ws.cell(row=2, column=3).value, "상한초과")   # critical
        self.assertEqual(ws.cell(row=2, column=6).value, -6.13)        # 면적%
        self.assertEqual(ws.cell(row=2, column=10).value, -5.76)       # 좌변%
        self.assertEqual(ws.cell(row=2, column=9).value, "정합")       # 정합성

    def test_summary_fields(self):
        ws = self._wb()["요약"]
        keys = {ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)}
        for expect in ("면적 확대율(%)", "기준 상한(%)", "상한 초과(건)"):
            self.assertIn(expect, keys)

    def test_overlap_sheet_has_image(self):
        wb = self._wb()
        self.assertGreaterEqual(len(wb["오버랩이미지"]._images), 1)

    def test_no_old_bbox_columns(self):
        """예전 bbox 컬럼(가로/세로 확대율·사유) 제거 확인."""
        ws = self._wb()["조각별상세"]
        header = [ws.cell(row=1, column=c).value for c in range(1, 15)]
        self.assertNotIn("가로 확대율(%)", header)
        self.assertNotIn("사유", header)


if __name__ == "__main__":
    unittest.main(verbosity=2)
